"""
UniSIRAJ CIA Helpdesk Agent - Microsoft Excel Integration Service
"""

import os
import io
import re
from datetime import datetime, date, timezone
from typing import Dict, Any, List, Optional, Tuple
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.student import Student, StudentStatus
from app.services.student_service import calculate_days_remaining, determine_student_status, get_current_business_date, clean_phone_number
from app.services.spreadsheet_service import parse_date_safely, normalize_phone_number
from app.services.audit_service import AuditService
from app.schemas.excel import (
    ExcelStatusResponse,
    ExcelStudentItem,
    ExcelFetchResponse,
    ExcelSearchResponse,
    ExcelAnalysisResponse,
    ExcelSyncResponse,
    DiscrepancyItem,
)
from app.core.logging import logger

DATA_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data")
DEFAULT_EXCEL_PATH = os.path.join(DATA_DIR, "unisiraj_students_master.xlsx")

COLUMN_SYNONYMS = {
    "matric_id": ["matric id", "matric_id", "matric number", "student id", "student_id", "id", "no matrik"],
    "student_name": ["student name", "student_name", "full name", "name", "nama", "student"],
    "nationality": ["nationality", "country", "citizenship", "warganegara", "negara"],
    "passport_number": ["passport number", "passport no", "passport", "no passport", "no pasport"],
    "visa_expiry_date": ["visa expiry date", "visa expiry", "expiry date", "tarikh tamat visa", "expiry"],
    "whatsapp_number": ["whatsapp number", "phone number", "phone", "contact number", "mobile", "no tel", "whatsapp"],
    "faculty": ["faculty", "kulliyyah", "school", "fakulti"],
    "program": ["program", "programme", "course", "degree", "kursus"],
    "emgs_status": ["emgs status", "emgs stage", "emgs", "status emgs"],
    "notes": ["notes", "remarks", "catatan", "comment"],
}


class ExcelService:
    DEFAULT_EXCEL_PATH = DEFAULT_EXCEL_PATH

    @classmethod
    def resolve_path(cls, filepath: Optional[str] = None) -> str:
        if filepath and filepath.strip():
            return filepath.strip()
        os.makedirs(DATA_DIR, exist_ok=True)
        return cls.DEFAULT_EXCEL_PATH

    @classmethod
    def ensure_master_excel_file(cls, filepath: Optional[str] = None) -> str:
        target = cls.resolve_path(filepath)
        if os.path.exists(target) and os.path.getsize(target) > 0:
            return target

        os.makedirs(os.path.dirname(target), exist_ok=True)
        logger.info(f"Generating initial master Microsoft Excel spreadsheet at: {target}")

        initial_records = [
            ("UIS-2023-INT084", "Tariq Al-Mansoor", "Yemen", "09482103", "2026-10-06", "+601123456789", "Faculty of Islamic Economics & Business", "Bachelor of Islamic Finance (Hons)", "In Progress", "Urgent passport renewal recommended"),
            ("UIS-2022-INT192", "Fatima Al-Zahra", "Indonesia", "ID8492019", "2026-10-25", "+601234567890", "Faculty of Shariah & Law", "Master of Comparative Law", "Approved", "Biometric verification completed"),
            ("UIS-2023-INT045", "Amina Zainab Bello", "Nigeria", "A09283411", "2026-11-02", "+601245678901", "Faculty of Computing & Informatics", "Bachelor of Computer Science", "Requires Biometrics", "Awaiting Val extension approval"),
            ("UIS-2024-INT301", "Muhammad Bilal Khan", "Pakistan", "PK9102834", "2026-11-18", "+601356789012", "Faculty of Islamic Economics & Business", "Bachelor of Halal Industry Management", "In Progress", "Document review queued"),
            ("UIS-2022-INT012", "Nour Al-Huda Mansouri", "Syria", "N7839201", "2026-12-05", "+601467890123", "Faculty of Quranic & Sunnah Studies", "PhD in Quranic Exegesis", "In Progress", "Tier 2 notification delivered"),
            ("UIS-2023-INT119", "Khaled Omar Al-Sayed", "Egypt", "EG4829103", "2026-12-14", "+601578901234", "Faculty of Shariah & Law", "Bachelor of Islamic Jurisprudence", "Approved", "Sticker endorsement pending"),
            ("UIS-2024-INT008", "Zayd Ibrahim Toure", "Mali", "ML2938471", "2026-12-28", "+601689012345", "Faculty of Education & Human Sciences", "Master of Islamic Education", "Not Started", "Student advised on fee settlement"),
            ("UIS-2021-INT099", "Maryam Hassan Dahir", "Somalia", "SO9384712", "2027-01-10", "+601790123456", "Faculty of Computing & Informatics", "Master of Data Science", "Approved", "Compliance milestone cleared"),
            ("UIS-2023-INT204", "Rashid Ahmed Chowdhury", "Bangladesh", "BD5839201", "2027-01-15", "+601801234567", "Faculty of Islamic Economics & Business", "Bachelor of Islamic Banking", "In Progress", "Tier 1 advisory dispatched"),
            ("UIS-2024-INT155", "Siti Nurul Hidayah", "Brunei", "BN4829104", "2027-02-01", "+601912345678", "Faculty of Shariah & Law", "Bachelor of Fiqh & Fatwa", "Not Started", "Initial checklist reviewed"),
            ("UIS-2023-INT067", "Hamza Mansour Al-Houthi", "Yemen", "YE982341", "2026-10-14", "+601145678901", "Faculty of Shariah & Law", "Bachelor of Fiqh & Usul (Hons)", "Requires Biometrics", "High-priority embassy attestation"),
            ("UIS-2023-INT502", "Zulaikha binti Abdul Rahman", "Indonesia", "ID829104", "2026-11-20", "+601789012345", "Faculty of Islamic Economics & Business", "Bachelor of Islamic Finance", "In Progress", "Renewal follow-up scheduled"),
            ("UIS-2024-INT503", "Ibrahim Musa Danjuma", "Nigeria", "NG482910", "2026-12-19", "+601890123456", "Faculty of Computing & Informatics", "Bachelor of Computer Science", "Not Started", "Notice delivered via WhatsApp"),
            ("UIS-2024-INT504", "Aysha Siddiqua", "Bangladesh", "BD729104", "2027-02-15", "+601401234567", "Faculty of Quranic & Islamic Studies", "Master of Sunnah Studies", "Approved", "Valid student pass in place"),
        ]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "UniSIRAJ_International_Students"

        # Apply professional corporate header styling
        headers = [
            "Matric ID", "Student Name", "Nationality", "Passport Number",
            "Visa Expiry Date", "WhatsApp Number", "Faculty", "Program",
            "EMGS Status", "Notes"
        ]

        header_fill = PatternFill(start_color="0D2040", end_color="0D2040", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style='thin', color='D0D5DD'),
            right=Side(style='thin', color='D0D5DD'),
            top=Side(style='thin', color='D0D5DD'),
            bottom=Side(style='thin', color='D0D5DD')
        )

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_data in initial_records:
            ws.append(list(row_data))

        # Format data rows
        for row in ws.iter_rows(min_row=2, max_row=len(initial_records) + 1, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = border_thin
                cell.font = Font(name="Calibri", size=10)

        # Set auto-width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(target)
        logger.info(f"Master Microsoft Excel spreadsheet created with {len(initial_records)} initial records.")
        return target

    @classmethod
    def _read_excel_dataframe(cls, filepath: Optional[str] = None, sheet_name: Optional[str] = None) -> Tuple[pd.DataFrame, str]:
        path = cls.ensure_master_excel_file(filepath)
        xl = pd.ExcelFile(path)
        sheet = sheet_name if (sheet_name and sheet_name in xl.sheet_names) else xl.sheet_names[0]
        df = pd.read_excel(xl, sheet_name=sheet)
        return df, sheet

    @classmethod
    def _map_columns(cls, df: pd.DataFrame) -> Dict[str, str]:
        mapping = {}
        for col in df.columns:
            col_clean = str(col).strip().lower()
            for standard_key, synonyms in COLUMN_SYNONYMS.items():
                if standard_key not in mapping and any(syn in col_clean for syn in synonyms):
                    mapping[standard_key] = col
                    break
        return mapping

    @classmethod
    def get_status(cls, filepath: Optional[str] = None) -> ExcelStatusResponse:
        path = cls.ensure_master_excel_file(filepath)
        size = os.path.getsize(path)
        mtime = datetime.fromtimestamp(os.path.getmtime(path), tz=timezone.utc).isoformat()
        xl = pd.ExcelFile(path)
        sheet_names = xl.sheet_names
        active_sheet = sheet_names[0] if sheet_names else "Sheet1"
        df = pd.read_excel(xl, sheet_name=active_sheet)

        return ExcelStatusResponse(
            is_connected=True,
            filename=os.path.basename(path),
            filepath=path,
            file_size_bytes=size,
            sheet_names=sheet_names,
            active_sheet=active_sheet,
            total_rows=len(df),
            last_modified=mtime,
            last_sync_timestamp=mtime,
            sync_status="SYNCHRONIZED",
        )

    @classmethod
    def fetch_records(
        cls,
        sheet_name: Optional[str] = None,
        page: int = 1,
        page_size: int = 25,
        filepath: Optional[str] = None,
    ) -> ExcelFetchResponse:
        df, actual_sheet = cls._read_excel_dataframe(filepath, sheet_name)
        mapping = cls._map_columns(df)
        today = get_current_business_date()

        total = len(df)
        total_pages = max(1, (total + page_size - 1) // page_size)
        start_idx = (page - 1) * page_size
        end_idx = min(start_idx + page_size, total)

        page_df = df.iloc[start_idx:end_idx]
        items: List[ExcelStudentItem] = []

        for offset, (_, row) in enumerate(page_df.iterrows()):
            row_num = start_idx + offset + 2  # 1-based, accounting for header row
            matric = str(row.get(mapping.get("matric_id", ""), "")).strip()
            name = str(row.get(mapping.get("student_name", ""), "")).strip()

            exp_val = row.get(mapping.get("visa_expiry_date", ""))
            exp_date, _ = parse_date_safely(exp_val)
            days = calculate_days_remaining(exp_date, today)
            st_status = determine_student_status(days)

            phone_val = row.get(mapping.get("whatsapp_number", ""))
            norm_phone, _ = normalize_phone_number(phone_val)

            items.append(
                ExcelStudentItem(
                    row_index=row_num,
                    matric_id=matric,
                    student_name=name,
                    nationality=str(row.get(mapping.get("nationality", ""), "") or "").strip() or None,
                    passport_number=str(row.get(mapping.get("passport_number", ""), "") or "").strip() or None,
                    visa_expiry_date=exp_date.isoformat() if exp_date else None,
                    whatsapp_number=norm_phone or clean_phone_number(phone_val),
                    faculty=str(row.get(mapping.get("faculty", ""), "") or "").strip() or None,
                    program=str(row.get(mapping.get("program", ""), "") or "").strip() or None,
                    emgs_status=str(row.get(mapping.get("emgs_status", ""), "") or "").strip() or None,
                    days_remaining=days,
                    status=st_status,
                    notes=str(row.get(mapping.get("notes", ""), "") or "").strip() or None,
                )
            )

        return ExcelFetchResponse(
            sheet_name=actual_sheet,
            total_rows=total,
            page=page,
            page_size=page_size,
            total_pages=total_pages,
            data=items,
        )

    @classmethod
    def search_students(
        cls,
        query: str,
        sheet_name: Optional[str] = None,
        filepath: Optional[str] = None,
    ) -> ExcelSearchResponse:
        df, actual_sheet = cls._read_excel_dataframe(filepath, sheet_name)
        mapping = cls._map_columns(df)
        today = get_current_business_date()
        q_clean = query.strip().lower()
        stopwords = {"in", "excel", "file", "master", "spreadsheet", "the", "for", "student", "records", "record", "of", "please", "search", "find", "show", "a"}
        tokens = [w for w in re.findall(r"[\w\-]+", q_clean) if w not in stopwords]

        matches: List[ExcelStudentItem] = []

        for idx, row in df.iterrows():
            row_str = " ".join([str(v) for v in row.values if pd.notna(v)]).lower()
            is_match = False
            if q_clean in row_str:
                is_match = True
            elif tokens and any(t in row_str for t in tokens):
                is_match = True

            if is_match:
                matric = str(row.get(mapping.get("matric_id", ""), "")).strip()
                name = str(row.get(mapping.get("student_name", ""), "")).strip()
                exp_val = row.get(mapping.get("visa_expiry_date", ""))
                exp_date, _ = parse_date_safely(exp_val)
                days = calculate_days_remaining(exp_date, today)
                st_status = determine_student_status(days)
                phone_val = row.get(mapping.get("whatsapp_number", ""))
                norm_phone, _ = normalize_phone_number(phone_val)

                matches.append(
                    ExcelStudentItem(
                        row_index=idx + 2,
                        matric_id=matric,
                        student_name=name,
                        nationality=str(row.get(mapping.get("nationality", ""), "") or "").strip() or None,
                        passport_number=str(row.get(mapping.get("passport_number", ""), "") or "").strip() or None,
                        visa_expiry_date=exp_date.isoformat() if exp_date else None,
                        whatsapp_number=norm_phone or clean_phone_number(phone_val),
                        faculty=str(row.get(mapping.get("faculty", ""), "") or "").strip() or None,
                        program=str(row.get(mapping.get("program", ""), "") or "").strip() or None,
                        emgs_status=str(row.get(mapping.get("emgs_status", ""), "") or "").strip() or None,
                        days_remaining=days,
                        status=st_status,
                        notes=str(row.get(mapping.get("notes", ""), "") or "").strip() or None,
                    )
                )

        return ExcelSearchResponse(
            query=query,
            total_matches=len(matches),
            matches=matches,
        )

    @classmethod
    def analyze_spreadsheet(
        cls,
        db: Session,
        sheet_name: Optional[str] = None,
        filepath: Optional[str] = None,
    ) -> ExcelAnalysisResponse:
        df, actual_sheet = cls._read_excel_dataframe(filepath, sheet_name)
        mapping = cls._map_columns(df)
        today = get_current_business_date()

        # Database index of students
        db_students = {s.matric_id.strip(): s for s in db.query(Student).all()}

        quality_metrics = {
            "missing_dates": 0,
            "missing_phones": 0,
            "missing_passports": 0,
            "invalid_dates": 0,
        }
        status_dist = {
            "SAFE": 0,
            "ACTION_NEEDED": 0,
            "URGENT": 0,
            "CRITICAL": 0,
            "EXPIRED": 0,
            "REVIEW_REQUIRED": 0,
        }
        discrepancies: List[DiscrepancyItem] = []
        excel_matrics = set()

        for idx, row in df.iterrows():
            matric = str(row.get(mapping.get("matric_id", ""), "")).strip()
            name = str(row.get(mapping.get("student_name", ""), "")).strip()
            if not matric:
                continue

            excel_matrics.add(matric)

            # Expiry date parsing
            exp_val = row.get(mapping.get("visa_expiry_date", ""))
            exp_date, date_err = parse_date_safely(exp_val)
            if not exp_val or pd.isna(exp_val):
                quality_metrics["missing_dates"] += 1
            elif date_err:
                quality_metrics["invalid_dates"] += 1

            days = calculate_days_remaining(exp_date, today)
            st_status = determine_student_status(days)
            status_dist[st_status] = status_dist.get(st_status, 0) + 1

            # Phone check
            phone_val = row.get(mapping.get("whatsapp_number", ""))
            if pd.isna(phone_val) or not str(phone_val).strip():
                quality_metrics["missing_phones"] += 1

            # Passport check
            passport_val = row.get(mapping.get("passport_number", ""))
            if pd.isna(passport_val) or not str(passport_val).strip():
                quality_metrics["missing_passports"] += 1

            # Compare against DB
            if matric not in db_students:
                discrepancies.append(
                    DiscrepancyItem(
                        matric_id=matric,
                        student_name=name,
                        discrepancy_type="MISSING_IN_DB",
                        description=f"Student {name} ({matric}) is present in Excel but not enrolled in active system database.",
                        excel_value=name,
                        db_value=None,
                    )
                )
            else:
                db_st = db_students[matric]
                # Compare expiry date
                if exp_date and db_st.visa_expiry_date and exp_date != db_st.visa_expiry_date:
                    discrepancies.append(
                        DiscrepancyItem(
                            matric_id=matric,
                            student_name=name,
                            discrepancy_type="FIELD_MISMATCH",
                            field_name="visa_expiry_date",
                            excel_value=exp_date.isoformat(),
                            db_value=db_st.visa_expiry_date.isoformat(),
                            description=f"Expiry date mismatch: Excel lists {exp_date.isoformat()}, while Database has {db_st.visa_expiry_date.isoformat()}.",
                        )
                    )
                # Compare phone number
                excel_phone, _ = normalize_phone_number(phone_val)
                if excel_phone and db_st.whatsapp_number and excel_phone != db_st.whatsapp_number:
                    discrepancies.append(
                        DiscrepancyItem(
                            matric_id=matric,
                            student_name=name,
                            discrepancy_type="FIELD_MISMATCH",
                            field_name="whatsapp_number",
                            excel_value=excel_phone,
                            db_value=db_st.whatsapp_number,
                            description=f"Phone number mismatch: Excel lists {excel_phone}, Database has {db_st.whatsapp_number}.",
                        )
                    )

        # Check DB students missing from Excel
        for matric, db_st in db_students.items():
            if matric not in excel_matrics:
                discrepancies.append(
                    DiscrepancyItem(
                        matric_id=matric,
                        student_name=db_st.student_name,
                        discrepancy_type="MISSING_IN_EXCEL",
                        description=f"Student {db_st.student_name} ({matric}) exists in DB but is missing from the active Excel spreadsheet.",
                        excel_value=None,
                        db_value=db_st.student_name,
                    )
                )

        return ExcelAnalysisResponse(
            filename=os.path.basename(cls.resolve_path(filepath)),
            active_sheet=actual_sheet,
            total_records=len(df),
            quality_metrics=quality_metrics,
            visa_status_distribution=status_dist,
            discrepancies=discrepancies,
            total_discrepancies=len(discrepancies),
            analyzed_at=datetime.now(timezone.utc).isoformat(),
        )

    @classmethod
    def sync_excel_to_db(
        cls,
        db: Session,
        sheet_name: Optional[str] = None,
        filepath: Optional[str] = None,
        user_id: Optional[int] = None,
    ) -> ExcelSyncResponse:
        df, _ = cls._read_excel_dataframe(filepath, sheet_name)
        mapping = cls._map_columns(df)
        today = get_current_business_date()

        records_created = 0
        records_updated = 0
        conflicts_resolved = 0

        for _, row in df.iterrows():
            matric = str(row.get(mapping.get("matric_id", ""), "")).strip()
            name = str(row.get(mapping.get("student_name", ""), "")).strip()
            if not matric or not name:
                continue

            exp_val = row.get(mapping.get("visa_expiry_date", ""))
            exp_date, _ = parse_date_safely(exp_val)
            days = calculate_days_remaining(exp_date, today)
            st_status = determine_student_status(days)

            phone_val = row.get(mapping.get("whatsapp_number", ""))
            norm_phone, _ = normalize_phone_number(phone_val)

            existing = db.query(Student).filter(Student.matric_id == matric).first()
            if existing:
                had_conflict = False
                if exp_date and existing.visa_expiry_date != exp_date:
                    existing.visa_expiry_date = exp_date
                    had_conflict = True
                if norm_phone and existing.whatsapp_number != norm_phone:
                    existing.whatsapp_number = norm_phone
                    had_conflict = True
                if name and existing.student_name != name:
                    existing.student_name = name

                faculty_val = str(row.get(mapping.get("faculty", ""), "") or "").strip()
                if faculty_val:
                    existing.faculty = faculty_val
                program_val = str(row.get(mapping.get("program", ""), "") or "").strip()
                if program_val:
                    existing.program = program_val
                passport_val = str(row.get(mapping.get("passport_number", ""), "") or "").strip()
                if passport_val:
                    existing.passport_number = passport_val

                existing.status = st_status
                existing.last_verified_at = datetime.now(timezone.utc)
                records_updated += 1
                if had_conflict:
                    conflicts_resolved += 1
            else:
                new_st = Student(
                    matric_id=matric,
                    student_name=name,
                    nationality=str(row.get(mapping.get("nationality", ""), "") or "").strip() or None,
                    passport_number=str(row.get(mapping.get("passport_number", ""), "") or "").strip() or None,
                    visa_expiry_date=exp_date,
                    whatsapp_number=norm_phone or clean_phone_number(phone_val),
                    faculty=str(row.get(mapping.get("faculty", ""), "") or "").strip() or None,
                    program=str(row.get(mapping.get("program", ""), "") or "").strip() or None,
                    status=st_status,
                    last_verified_at=datetime.now(timezone.utc),
                )
                db.add(new_st)
                records_created += 1

        db.commit()

        AuditService.log_event(
            db=db,
            action="EXCEL_SYNC_TO_DB",
            resource_type="EXCEL",
            user_id=user_id,
            details={
                "records_created": records_created,
                "records_updated": records_updated,
                "conflicts_resolved": conflicts_resolved,
            },
        )

        return ExcelSyncResponse(
            direction="excel_to_db",
            records_created=records_created,
            records_updated=records_updated,
            conflicts_resolved=conflicts_resolved,
            total_processed=records_created + records_updated,
            message=f"Synchronized Excel to Database: {records_created} created, {records_updated} updated ({conflicts_resolved} conflicts resolved).",
            synced_at=datetime.now(timezone.utc).isoformat(),
        )

    @classmethod
    def sync_db_to_excel(
        cls,
        db: Session,
        sheet_name: Optional[str] = None,
        filepath: Optional[str] = None,
        user_id: Optional[int] = None,
    ) -> ExcelSyncResponse:
        target = cls.resolve_path(filepath)
        students = db.query(Student).order_by(Student.matric_id).all()
        today = get_current_business_date()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name or "UniSIRAJ_International_Students"

        headers = [
            "Matric ID", "Student Name", "Nationality", "Passport Number",
            "Visa Expiry Date", "WhatsApp Number", "Faculty", "Program",
            "EMGS Status", "Days Remaining", "Compliance Status", "Last Synced"
        ]

        header_fill = PatternFill(start_color="0D2040", end_color="0D2040", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style='thin', color='D0D5DD'),
            right=Side(style='thin', color='D0D5DD'),
            top=Side(style='thin', color='D0D5DD'),
            bottom=Side(style='thin', color='D0D5DD')
        )

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        sync_time_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

        for s in students:
            days = calculate_days_remaining(s.visa_expiry_date, today)
            st_status = determine_student_status(days)

            row = [
                s.matric_id,
                s.student_name,
                s.nationality or "",
                s.passport_number or "",
                s.visa_expiry_date.isoformat() if s.visa_expiry_date else "",
                s.whatsapp_number or "",
                s.faculty or "",
                s.program or "",
                "Approved",
                days if days is not None else "N/A",
                st_status,
                sync_time_str,
            ]
            ws.append(row)

        for row in ws.iter_rows(min_row=2, max_row=len(students) + 1, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = border_thin
                cell.font = Font(name="Calibri", size=10)

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(target)

        AuditService.log_event(
            db=db,
            action="EXCEL_SYNC_DB_TO_EXCEL",
            resource_type="EXCEL",
            user_id=user_id,
            details={"students_exported": len(students)},
        )

        return ExcelSyncResponse(
            direction="db_to_excel",
            records_created=0,
            records_updated=len(students),
            conflicts_resolved=0,
            total_processed=len(students),
            message=f"Successfully exported {len(students)} active database records to Microsoft Excel master workbook.",
            synced_at=datetime.now(timezone.utc).isoformat(),
        )

    @classmethod
    def save_uploaded_file(cls, file_bytes: bytes, filename: str, filepath: Optional[str] = None) -> ExcelStatusResponse:
        target = cls.resolve_path(filepath)
        # Verify it's valid Excel by opening
        xl = pd.ExcelFile(io.BytesIO(file_bytes))
        with open(target, "wb") as f:
            f.write(file_bytes)

        return cls.get_status(target)

    @classmethod
    def get_excel_bytes(cls, filepath: Optional[str] = None) -> bytes:
        target = cls.ensure_master_excel_file(filepath)
        with open(target, "rb") as f:
            return f.read()
